# -*- coding: utf-8 -*-
import os, sys, cv2, csv, json, gc
import numpy as np
import fitz
import pytesseract
from openpyxl import Workbook
from PIL import Image

try:
    from docx import Document
except ImportError:
    Document = None

from utils import (
    auto_adjust_excel_column_width,
    merge_2d_arrays_horizontally,
    sanitize_excel_text
)

# ==============================
# ローカルOCR抽出タスク (Tesseract)
# ==============================
def check_tesseract_installation():
    if sys.platform.startswith("win"):
        tess_path = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        if os.path.exists(tess_path): pytesseract.pytesseract.tesseract_cmd = tess_path
    try: pytesseract.get_tesseract_version()
    except Exception: raise Exception("Tesseract OCRがインストールされていないか、PATHが通っていません。")

def extract_tesseract_task(files, save_dir, options, ui):
    check_tesseract_installation()
    files = [f for f in files if f.lower().endswith(".pdf")]
    if not files: raise Exception("PDFファイルが含まれていません。")
    out_format = options.get("out_format", "xlsx")
    crop_regions = options.get("crop_regions", [])
    
    for i, f in enumerate(files, 1):
        if ui.is_cancelled(): return
        ui.update_overall(i, len(files), f"全体の進捗 ( {i} / {len(files)} ファイル )")
        base = os.path.splitext(os.path.basename(f))[0]
        doc = fitz.open(f)
        digits = max(2, len(str(len(doc))))
        total_pages = len(doc)
        for page_num in range(total_pages):
            if ui.is_cancelled(): return
            ui.set_indeterminate(f"Tesseractで解析中... ( {page_num+1} / {total_pages} ページ )")
            pix = doc[page_num].get_pixmap(dpi=300)
            img_array = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.h, pix.w, pix.n)
            if pix.n == 4: img_array = cv2.cvtColor(img_array, cv2.COLOR_RGBA2RGB)
            elif pix.n == 1: img_array = cv2.cvtColor(img_array, cv2.COLOR_GRAY2RGB)
            
            cropped_images = []
            if crop_regions:
                h, w = img_array.shape[:2]
                for (rx1, ry1, rx2, ry2) in crop_regions:
                    cropped_images.append(img_array[int(min(ry1, ry2)*h):int(max(ry1, ry2)*h), int(min(rx1, rx2)*w):int(max(rx1, rx2)*w)])
            else: cropped_images.append(img_array)
                
            all_regions_data = []
            for crop_img in cropped_images:
                try:
                    text = pytesseract.image_to_string(Image.fromarray(crop_img), lang="jpn+eng")
                    lines = [line.strip() for line in text.split('\n') if line.strip()]
                    if lines:
                        if (crop_regions and crop_img.shape[0] > crop_img.shape[1] * 1.5) or (len(lines) > 1 and all(len(l) <= 2 for l in lines)):
                            all_regions_data.append([["".join(lines)]])
                        else: all_regions_data.append([[l] for l in lines])
                    else: all_regions_data.append([[""]])
                except Exception as e: raise Exception(f"Tesseract OCRエラー: {e}")
                    
            merged_data = merge_2d_arrays_horizontally(all_regions_data)
            final_data = []
            page_info = f"{page_num+1}/{len(doc)}"
            if crop_regions:
                final_data.append(["ページ番号"] + [f"抽出範囲{idx+1}" for idx in range(len(cropped_images))])
            else:
                final_data.append(["ページ番号", "抽出テキスト"])
            for row in merged_data: final_data.append([page_info] + row)
            
            save_path = os.path.join(save_dir, f"{base}_Page_{str(page_num+1).zfill(digits)}_Tesseract抽出")
            if out_format == "xlsx":
                wb = Workbook(); ws = wb.active; ws.title = f"Page_{str(page_num+1).zfill(digits)}"
                for r_idx, row_data in enumerate(final_data, 1):
                    for c_idx, val in enumerate(row_data, 1): 
                        ws.cell(row=r_idx, column=c_idx, value=sanitize_excel_text(val))
                auto_adjust_excel_column_width(ws); wb.save(f"{save_path}.xlsx")
            elif out_format == "csv":
                with open(f"{save_path}.csv", "w", encoding="utf-8-sig", newline="") as f_out: csv.writer(f_out).writerows(final_data)
            elif out_format == "txt":
                with open(f"{save_path}.txt", "w", encoding="utf-8") as f_out:
                    for row_data in final_data: f_out.write("\t".join(row_data) + "\n")
            elif out_format == "json":
                with open(f"{save_path}.json", "w", encoding="utf-8") as f_out: json.dump(final_data, f_out, ensure_ascii=False, indent=2)
            elif out_format == "md":
                with open(f"{save_path}.md", "w", encoding="utf-8") as f_out:
                    if final_data:
                        f_out.write("| " + " | ".join(map(str, final_data[0])) + " |\n")
                        f_out.write("|" + "|".join(["---"] * len(final_data[0])) + "|\n")
                        for row in final_data[1:]: f_out.write("| " + " | ".join(map(lambda x: str(x).replace('\n', '<br>'), row)) + " |\n")
            elif out_format == "docx":
                if Document is None: raise Exception("python-docx ライブラリがインストールされていません。")
                doc_out = Document()
                if final_data:
                    table = doc_out.add_table(rows=len(final_data), cols=max(len(r) for r in final_data))
                    table.style = 'Table Grid'
                    for r_idx, row_data in enumerate(final_data):
                        row_cells = table.rows[r_idx].cells
                        for c_idx, val in enumerate(row_data):
                            if c_idx < len(row_cells):
                                row_cells[c_idx].text = str(val)
                doc_out.save(f"{save_path}.docx")
        doc.close(); gc.collect()
        ui.set_determinate(total_pages, total_pages, "完了")