# -*- coding: utf-8 -*-
import os, sys, cv2, csv, time, json, random, gc, re
import threading
import concurrent.futures
import numpy as np
import fitz
from openpyxl import Workbook
from PIL import Image
import google.generativeai as genai

try:
    from docx import Document
except ImportError:
    Document = None

from common import (
    auto_adjust_excel_column_width,
    parse_row_data,
    merge_2d_arrays_horizontally,
    sanitize_excel_text
)

from engines import expand_crop_rect_for_intersecting_objects

# ==============================
# Gemini API データ抽出タスク
# ==============================
def extract_gemini_task(files, save_dir, options, ui):
    files = [f for f in files if f.lower().endswith(".pdf")]
    if not files: raise Exception("PDFファイルが含まれていません。")
    api_key = options.get("api_key", "")
    genai.configure(api_key=api_key)
    
    models_to_try = options.get("models_to_try", ["gemini-2.5-flash"])
    out_format = options.get("out_format", "xlsx")
    crop_regions = options.get("crop_regions", [])
    extract_mode = options.get("extract_mode", "table")
    
    # JSON（表）形式で抽出するかどうかを判定
    is_table_format = out_format in ["csv", "xlsx", "json", "md", "docx"] and extract_mode == "table"

    if is_table_format:
        if crop_regions:
            # 複数領域を1枚の画像として送信するため、各領域のデータが独立した配列として返るようプロンプトを調整
            prompt = """
            あなたは優秀なデータ入力オペレーターです。添付された画像には、1つ以上のテキスト領域（抽出範囲）が上から下に向かって配置されています。
            それぞれの領域のテキストを読み取り、以下のJSON形式で出力してください。

            【特別ルール（絶対厳守）】
            - 画像内に複数の領域がある場合、上から順番にそれぞれのデータを読み取ってください。
            - 各領域内で意味的なデータの区切り（項目名と値、または別々の行など）がある場合は、改行で繋げずに、必ず配列（リスト）の別々の要素に分割（セル分け）して出力してください。
            - 縦書きのテキストは、必ず繋げて横書きの1つの文字列に変換してください。
            - 『〃』『...』『同上』などの同上を表す記号が記載されている場合は、空白扱いにせず、必ずその記号をそのまま出力してください。

            【出力形式（絶対厳守）】
            以下のように、画像内の領域の数と同じ要素数を持つ配列を `regions` キーの値として出力してください。
            各要素は、その領域から抽出された各データを格納した配列です。
            {
              "regions": [
                [ "領域1のデータ1", "領域1のデータ2", "領域1のデータ3..." ],
                [ "領域2のデータ1..." ]
              ]
            }
            """
        else:
            prompt = """
            あなたは優秀なデータ入力オペレーターです。添付された図面管理台帳などの表画像を読み取り、正確なJSONデータを作成してください。
            
            【データの分離精度を最優先する特別ルール（超重要）】
            - データの見た目や文脈の意味よりも、「表の縦の罫線」や「文字列間の大きな空白」などの【物理的な列の区切り】を絶対的な基準として最優先し、物理的に分かれているデータは必ず別の要素として分割してください。
            - 意味的に繋がっているように見えても、罫線や空白で区切られていれば絶対に1つの要素に結合しないでください。
            - データが存在しない「空白セル」の場合は無視せず、必ず `""` (空文字) を該当位置に挿入し、列の位置を厳密に揃えてください。ただし、『〃』『...』『同上』などの同上を表す記号が記載されている場合は、絶対に空白にせず、その記号をそのまま出力してください。
            - 1行のデータを丸ごと1つの文字列に繋げて出力することは絶対に禁止します。必ず各セルごとにリスト内で分割してください。
            - 縦書きのテキスト（文字が縦に並んでいるもの）は、1文字ずつ分割したり複数行に分けたりせず、必ず繋げて横書きの1つの文字列に変換して出力してください。

            【出力形式（絶対厳守）】
            シンプルな配列（リスト）で出力してください。
            {
              "header": ["列1名", "列2名", "列3名", ...],
              "rows": [
                ["データ1", "データ2", "", "データ4", ...],
                ["データ1", "データ2", "データ3", "", ...]
              ]
            }
            """
    else:
        if crop_regions:
            prompt = """
            添付された画像には複数の抽出領域が含まれています。各領域のテキストを読み取り、領域ごとに以下の形式でプレーンテキストとして出力してください。
            縦書きの文章は改行を取り除き、横書きに変換して出力してください。
            『〃』『...』『同上』などの同上を表す記号が記載されている場合は、必ずそのまま出力してください。
            ---領域1---
            テキスト内容
            ---領域2---
            テキスト内容
            """
        else:
            prompt = "この画像に記載されている手書きの文字や文章を可能な限り正確に読み取り、プレーンテキストとして出力してください。また、縦書きの文章は改行を取り除き、横書きに変換して出力してください。『〃』『...』『同上』などの同上を表す記号が記載されている場合は、必ずそのまま出力してください。"

    # オプションの取得
    api_plan = options.get("api_plan", "free")
    api_rpm = options.get("api_rpm", 12)
    temperature = options.get("temperature", 0.0)
    disable_safety = options.get("disable_safety", True)
    max_tokens = options.get("max_tokens", 8192)
    custom_prompt = options.get("custom_prompt", "")
    threads_setting = options.get("threads", 1 if api_plan == "free" else 5)
    
    if custom_prompt.strip():
        prompt += f"\n\n【ユーザーからの追加指示（厳守すること）】\n{custom_prompt.strip()}"

    generation_config_base = {
        "temperature": temperature,
        "max_output_tokens": max_tokens,
    }
    
    if is_table_format:
        generation_config = generation_config_base.copy()
        generation_config["response_mime_type"] = "application/json"
    else:
        generation_config = generation_config_base.copy()

    safety_settings = None
    if disable_safety:
        safety_settings = [
            {"category": "HARM_CATEGORY_HARASSMENT", "threshold": "BLOCK_NONE"},
            {"category": "HARM_CATEGORY_HATE_SPEECH", "threshold": "BLOCK_NONE"},
            {"category": "HARM_CATEGORY_SEXUALLY_EXPLICIT", "threshold": "BLOCK_NONE"},
            {"category": "HARM_CATEGORY_DANGEROUS_CONTENT", "threshold": "BLOCK_NONE"}
        ]

    MAX_RPM = api_rpm if api_rpm > 0 else 1
    max_workers = max(1, threads_setting)

    # グローバルな状態管理
    request_timestamps = []
    rate_limit_lock = threading.Lock()
    progress_lock = threading.Lock()
    completed_pages = 0
    shared_context = {"fatal_error": None}

    # 全ファイル・全ページのタスクリストを事前に作成
    page_tasks = []
    ui.set_indeterminate("処理対象のページをスキャン中...")
    for f in files:
        if ui.is_cancelled(): return
        try:
            doc = fitz.open(f)
            total_pages = len(doc)
            doc.close()
            for page_num in range(total_pages):
                page_tasks.append({
                    "file_path": f,
                    "page_num": page_num,
                    "total_pages": total_pages
                })
        except Exception as e:
            print(f"Failed to scan {f}: {e}")
            continue

    total_tasks = len(page_tasks)
    if total_tasks == 0:
        return
        
    ui.update_overall(0, total_tasks, f"全体の進捗 ( 0 / {total_tasks} ページ完了 )")

    def process_single_page_task(task_info):
        """1ページ分の全領域を1回のAPIリクエストで処理するタスク"""
        if shared_context["fatal_error"]: return
        if ui.is_cancelled(): return
        
        f_path = task_info["file_path"]
        page_num = task_info["page_num"]
        total_p = task_info["total_pages"]
        base = os.path.splitext(os.path.basename(f_path))[0]
        digits = max(2, len(str(total_p)))
        
        try:
            # 1. ページ画像の抽出
            doc = fitz.open(f_path)
            pix = doc[page_num].get_pixmap(dpi=300)
            doc.close()
            
            img_array = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.h, pix.w, pix.n)
            if pix.n == 4: img_array = cv2.cvtColor(img_array, cv2.COLOR_RGBA2RGB)
            elif pix.n == 1: img_array = cv2.cvtColor(img_array, cv2.COLOR_GRAY2RGB)

            # 複数の領域を1つの画像に結合するロジック
            combined_img = None
            cropped_info = [] # 結合後のパース用メタデータ
            
            if crop_regions:
                h_img, w_img = img_array.shape[:2]
                cropped_images = []
                for region in crop_regions:
                    rx1, ry1, rx2, ry2 = region[:4]
                    is_vert = region[4] if len(region) > 4 else False
                    is_line = is_vert or abs(ry2 - ry1) < 0.03 or abs(rx2 - rx1) < 0.03
                    
                    if not is_table_format or is_line:
                        x1, y1, x2, y2 = expand_crop_rect_for_intersecting_objects(img_array, rx1, ry1, rx2, ry2)
                    else:
                        x1, y1 = int(min(rx1, rx2) * w_img), int(min(ry1, ry2) * h_img)
                        x2, y2 = int(max(rx1, rx2) * w_img), int(max(ry1, ry2) * h_img)
                        
                    crop = img_array[y1:y2, x1:x2]
                    cropped_images.append(crop)
                    cropped_info.append(crop.shape)
                
                # 画像を縦に連結（間に少しの余白を入れる）
                if cropped_images:
                    max_w = max(img.shape[1] for img in cropped_images)
                    total_h = sum(img.shape[0] for img in cropped_images) + (len(cropped_images) - 1) * 20
                    
                    combined_array = np.full((total_h, max_w, 3), 255, dtype=np.uint8) # 白背景
                    
                    current_y = 0
                    for img in cropped_images:
                        h_crop, w_crop = img.shape[:2]
                        combined_array[current_y:current_y+h_crop, 0:w_crop] = img
                        current_y += h_crop + 20
                        
                    combined_img = Image.fromarray(combined_array)
            else:
                combined_img = Image.fromarray(img_array)
                cropped_info = [img_array.shape]
            
            # 画像リサイズ処理
            max_size = 2048
            if max(combined_img.width, combined_img.height) > max_size:
                ratio = max_size / max(combined_img.width, combined_img.height)
                new_size = (int(combined_img.width * ratio), int(combined_img.height * ratio))
                combined_img = combined_img.resize(new_size, Image.Resampling.LANCZOS)
            
            # 2. 1ページにつき1回のAPIリクエスト
            max_retries = 8
            extracted_text, success, last_error = "", False, ""
            
            for attempt in range(max_retries):
                if shared_context["fatal_error"]: raise Exception(shared_context["fatal_error"])
                if ui.is_cancelled(): return
                required_sleep = 0.0
                is_fatal_quota = False
                
                for model_name in models_to_try:
                    if shared_context["fatal_error"]: raise Exception(shared_context["fatal_error"])
                    if ui.is_cancelled(): return
                    total_wait = 0.0
                    
                    with rate_limit_lock:
                        now = time.time()
                        nonlocal request_timestamps
                        request_timestamps = [t for t in request_timestamps if now - t < 60.0]
                        min_interval = 1.5
                        
                        if len(request_timestamps) >= MAX_RPM:
                            expected_run_time = request_timestamps[0] + 60.0
                            total_wait = expected_run_time - now
                        else:
                            if request_timestamps:
                                last_planned_time = request_timestamps[-1]
                                if now < last_planned_time + min_interval:
                                    expected_run_time = last_planned_time + min_interval
                                    total_wait = expected_run_time - now
                                else:
                                    expected_run_time = now
                                    total_wait = 0.0
                            else:
                                expected_run_time = now
                                total_wait = 0.0
                        request_timestamps.append(expected_run_time)

                    if total_wait > 0:
                        if total_wait > 2.0:
                            ui.set_indeterminate(f"通信間隔調整中... (約 {int(total_wait)}秒待機)")
                        wait_step = 0.5
                        while total_wait > 0:
                            if shared_context["fatal_error"]: raise Exception(shared_context["fatal_error"])
                            if ui.is_cancelled(): return
                            time.sleep(min(wait_step, total_wait))
                            total_wait -= wait_step
                            
                    ui.set_indeterminate(f"AI解析中... ( P.{page_num+1} )")
                    
                    try:
                        model = genai.GenerativeModel(model_name)
                        contents = [prompt, combined_img]
                        if generation_config:
                            response = model.generate_content(contents, generation_config=generation_config, safety_settings=safety_settings)
                        else:
                            response = model.generate_content(contents, safety_settings=safety_settings)
                            
                        if not response.parts: raise Exception("安全フィルタ等によりブロックされました。")
                        extracted_text = normalize_text(response.text.strip())
                        success = True
                        break
                    except Exception as api_err: 
                        err_str = str(api_err)
                        last_error = err_str
                        if "429" in err_str or "Quota" in err_str:
                            m = re.search(r'retry in ([\d\.]+)s', err_str, re.IGNORECASE | re.DOTALL)
                            if not m: m = re.search(r'seconds:\s*(\d+)', err_str, re.IGNORECASE | re.DOTALL)
                            if m:
                                wait_sec = float(m.group(1))
                                if wait_sec > 15.0:  # 長い待機時間を要求された場合は即座に打ち切る
                                    msg = f"APIの利用枠（バースト制限）を超過しました（約{int(wait_sec)}秒の待機を要求されました）。\n処理全体を即時中断します。しばらく時間をおいてから再試行してください。"
                                    shared_context["fatal_error"] = msg
                                    raise Exception(msg)
                                required_sleep = wait_sec + 2.0 
                            else:
                                if "per day" in err_str.lower() or "perday" in err_str.lower():
                                    msg = "1日のAPI利用上限に到達しました。\n無料枠の場合は明日以降に再度お試しください。"
                                    shared_context["fatal_error"] = msg
                                    raise Exception(msg)
                                else:
                                    if attempt >= 2:
                                        msg = "APIの利用枠（クォータ）を超過している可能性が高いです。\n回復しないため処理全体を即時中断します。プランや利用状況を確認してください。"
                                        shared_context["fatal_error"] = msg
                                        raise Exception(msg)
                                    required_sleep = 4.0 + random.uniform(1.0, 3.0) 
                        elif "404" in err_str: 
                            msg = f"モデルが存在しないか、利用する権限がありません。\n詳細: {err_str}"
                            shared_context["fatal_error"] = msg
                            raise Exception(msg)
                        continue
                
                if success or is_fatal_quota: break
                if shared_context["fatal_error"]: raise Exception(shared_context["fatal_error"])
                if ui.is_cancelled(): return
                
                if required_sleep > 0: sleep_time = required_sleep
                else: sleep_time = min(60.0, 4.0 * (2 ** attempt)) + random.uniform(1.0, 3.0)
                    
                wait_step = 0.5
                while sleep_time > 0:
                    if shared_context["fatal_error"]: raise Exception(shared_context["fatal_error"])
                    if ui.is_cancelled(): return
                    time.sleep(min(wait_step, sleep_time))
                    sleep_time -= wait_step
                    
            # 3. 抽出データのパースと各領域への分配
            all_regions_data = []
            if success:
                if is_table_format:
                    try:
                        clean_text = extracted_text.strip()
                        json_match = re.search(r'\{.*\}', clean_text, re.DOTALL)
                        if json_match: clean_text = json_match.group(0)
                        
                        # データ量超過で途切れたJSONを強制的に修復する堅牢なロジック
                        def _robust_json_parse(json_str):
                            try:
                                return json.loads(json_str)
                            except json.JSONDecodeError as initial_err:
                                try:
                                    s = json_str.strip()
                                    quote_count = len(re.findall(r'(?<!\\)"', s))
                                    if quote_count % 2 != 0: s += '"'
                                    s = re.sub(r',\s*$', '', s) # 末尾のカンマを削除
                                    open_brackets = s.count('[') - s.count(']')
                                    open_braces = s.count('{') - s.count('}')
                                    if open_brackets > 0: s += ']' * open_brackets
                                    if open_braces > 0: s += '}' * open_braces
                                    return json.loads(s)
                                except Exception:
                                    try:
                                        # カンマ以降の不完全な文字列を丸ごと削ってから閉じる
                                        s2 = json_str.strip()
                                        last_comma = s2.rfind(',')
                                        if last_comma != -1:
                                            s2 = s2[:last_comma]
                                            quote_count = len(re.findall(r'(?<!\\)"', s2))
                                            if quote_count % 2 != 0: s2 += '"'
                                            ob = s2.count('[') - s2.count(']')
                                            oc = s2.count('{') - s2.count('}')
                                            if ob > 0: s2 += ']' * ob
                                            if oc > 0: s2 += '}' * oc
                                            return json.loads(s2)
                                    except Exception:
                                        pass
                                raise initial_err

                        data = _robust_json_parse(clean_text)
                        
                        if crop_regions:
                            # { "regions": [ [data1], [data2] ] } の形式を想定
                            parsed_regions = data.get("regions", [])
                            if not isinstance(parsed_regions, list):
                                parsed_regions = [parsed_regions]
                                
                            # 抽出した領域数と設定された領域数が合わない場合のフォールバック
                            for i in range(len(crop_regions)):
                                if i < len(parsed_regions):
                                    rows = parsed_regions[i]
                                    if not isinstance(rows, list): rows = [rows]
                                else:
                                    rows = [f"領域{i+1}のデータ取得失敗"]
                                    
                                page_data_to_write = []
                                h_crop, w_crop = cropped_info[i][:2]
                                clean_rows = []
                                for r in rows:
                                    val = " ".join([str(x) for x in r]) if isinstance(r, list) else str(r)
                                    if '\n' in val:
                                        lines = [l.strip() for l in val.split('\n') if l.strip()]
                                        if all(len(l) <= 2 for l in lines): val = "".join(lines)
                                    clean_rows.append(val)
                                    
                                if h_crop > w_crop * 1.5 and all(len(x.strip()) <= 2 for x in clean_rows if x.strip()): 
                                    page_data_to_write.append(["".join(clean_rows)])
                                else:
                                    # 改行で結合せず、そのままの配列を1行のデータとして扱う（セルごとに横並びに分割）
                                    if clean_rows:
                                        page_data_to_write.append(clean_rows)
                                    else:
                                        page_data_to_write.append([""])
                                        
                                all_regions_data.append(page_data_to_write)
                        else:
                            header = data.get("header", [])
                            rows = data.get("rows", [])
                            if not header and not rows and isinstance(data, list):
                                if data: header, rows = (data[0] if isinstance(data[0], list) else [str(data[0])]), data[1:]
                            safe_header = [str(x).strip() for x in header] if isinstance(header, list) else []
                            page_col_count = len(safe_header)
                            for r in rows:
                                if isinstance(r, list) and len(r) > page_col_count: page_col_count = len(r)
                            if not safe_header: safe_header = [f"列{idx+1}" for idx in range(page_col_count)]
                            page_data_to_write = []
                            padded_header = (safe_header + [""] * page_col_count)[:page_col_count]
                            page_data_to_write.append(padded_header)
                            for row_data in rows:
                                parsed_r = parse_row_data(row_data)
                                safe_row_local = []
                                for val in (parsed_r + [""] * page_col_count)[:page_col_count]:
                                    v_str = str(val)
                                    if '\n' in v_str:
                                        lines = [l.strip() for l in v_str.split('\n') if l.strip()]
                                        if len(lines) > 1 and all(len(l) <= 2 for l in lines): v_str = "".join(lines)
                                    safe_row_local.append(v_str)
                                if any(v != "" for v in safe_row_local): page_data_to_write.append(safe_row_local)
                            all_regions_data.append(page_data_to_write)
                    except Exception as e:
                        # パース失敗時は、領域ごとにエラーをセット
                        for _ in range(len(crop_regions) if crop_regions else 1):
                            all_regions_data.append([[f"JSONパースエラー (データ量超過の可能性): {e}"]])
                else:
                    if crop_regions:
                        # テキストフォーマットでの複数領域パース（"---領域X---"で分割）
                        raw_parts = re.split(r'---領域\d+---', extracted_text.strip())
                        parts = [p.strip() for p in raw_parts if p.strip()]
                        for i in range(len(crop_regions)):
                            if i < len(parts):
                                val = parts[i].strip()
                                if val:
                                    lines = [l.strip() for l in val.split('\n') if l.strip()]
                                    all_regions_data.append([lines] if lines else [[""]])
                                else:
                                    all_regions_data.append([[""]])
                            else:
                                all_regions_data.append([[""]])
                    else:
                        lines = [line.strip() for line in extracted_text.strip().split('\n') if line.strip()]
                        all_regions_data.append([lines] if lines else [[""]])
            else:
                for _ in range(len(crop_regions) if crop_regions else 1):
                    all_regions_data.append([[f"AI抽出失敗: {last_error}"]])
                    
            # 4. ページデータの統合とファイル保存
            if not is_table_format:
                merged_row = []
                for region_data in all_regions_data:
                    for r in region_data:
                        for c in r:
                            if str(c).strip(): merged_row.append(str(c).strip())
                merged_data = [merged_row] if merged_row else [[""]]
            else:
                merged_data = merge_2d_arrays_horizontally(all_regions_data)
                
            final_data = []
            page_info_str = f"{page_num+1}/{total_p}"
            
            if not is_table_format:
                max_cols = max((len(r) for r in merged_data), default=1)
                header = ["ページ番号"] + [f"テキスト{i+1}" for i in range(max_cols)]
                final_data.append(header)
                for row in merged_data: final_data.append([page_info_str] + row)
            elif crop_regions:
                header = ["ページ番号"]
                if all_regions_data and merged_data:
                    # 分割されたセル数に合わせてヘッダーを動的に生成する
                    for i, region_data in enumerate(all_regions_data):
                        num_cols = len(region_data[0]) if region_data and len(region_data) > 0 else 1
                        if num_cols == 1:
                            header.append(f"抽出範囲{i+1}")
                        else:
                            for j in range(num_cols):
                                header.append(f"抽出範囲{i+1}-{j+1}")
                else:
                    header = ["ページ番号"] + [f"抽出範囲{idx+1}" for idx in range(len(crop_regions))]
                
                final_data.append(header)
                for row in merged_data: final_data.append([page_info_str] + row)
            else:
                for r_idx, row in enumerate(merged_data):
                    if r_idx == 0: final_data.append(["ページ番号"] + row)
                    else: final_data.append([page_info_str] + row)
            
            # 元のファイル名（_Page_XX_AI抽出）のみを使用する（タイムスタンプは付与しない）
            save_path = os.path.join(save_dir, f"{base}_Page_{str(page_num+1).zfill(digits)}_AI抽出")
            if out_format == "xlsx":
                wb = Workbook(); ws = wb.active; ws.title = f"Page_{str(page_num+1).zfill(digits)}"
                for r_idx, row_data in enumerate(final_data, 1):
                    for c_idx, val in enumerate(row_data, 1): ws.cell(row=r_idx, column=c_idx, value=sanitize_excel_text(val))
                auto_adjust_excel_column_width(ws); wb.save(f"{save_path}.xlsx")
            elif out_format == "csv":
                with open(f"{save_path}.csv", "w", encoding="utf-8-sig", newline="") as f_out: csv.writer(f_out).writerows(final_data)
            elif out_format == "txt":
                with open(f"{save_path}.txt", "w", encoding="utf-8") as f_out:
                    for row_data in final_data: f_out.write("\t".join(row_data) + "\n")
            elif out_format == "json":
                with open(f"{save_path}.json", "w", encoding="utf-8") as f_out: json.dump(final_data, f_out, ensure_ascii=False, indent=2)
            elif out_format == "md":
                with open(f"{save_path}.md", "w", encoding="utf-8") as f_out:
                    if final_data:
                        f_out.write("| " + " | ".join(map(str, final_data[0])) + " |\n")
                        f_out.write("|" + "|".join(["---"] * len(final_data[0])) + "|\n")
                        for row in final_data[1:]: f_out.write("| " + " | ".join(map(lambda x: str(x).replace('\n', '<br>'), row)) + " |\n")
            elif out_format == "docx":
                if Document is None: raise Exception("python-docx ライブラリがインストールされていません。")
                doc_out = Document()
                if final_data:
                    table = doc_out.add_table(rows=len(final_data), cols=max(len(r) for r in final_data))
                    table.style = 'Table Grid'
                    for r_idx, row_data in enumerate(final_data):
                        row_cells = table.rows[r_idx].cells
                        for c_idx, val in enumerate(row_data):
                            if c_idx < len(row_cells):
                                row_cells[c_idx].text = str(val)
                doc_out.save(f"{save_path}.docx")
                
            # メモリ解放
            gc.collect()

        except Exception as e:
            if shared_context["fatal_error"]:
                raise Exception(shared_context["fatal_error"])
            raise Exception(f"ページ {page_num+1} の処理中にエラーが発生しました: {e}")

    # 完全にフラット化されたタスクリストをスレッドプールで一気に並列処理
    with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = [executor.submit(process_single_page_task, task) for task in page_tasks]
        
        for future in concurrent.futures.as_completed(futures):
            # 共有エラーフラグが立っている、またはキャンセルされた場合は全体をシャットダウン
            if ui.is_cancelled() or shared_context["fatal_error"]:
                executor.shutdown(wait=False, cancel_futures=True)
                if shared_context["fatal_error"]:
                    raise Exception(shared_context["fatal_error"])
                return
            
            # タスク内で発生した例外（バースト制限など）をここでキャッチしてメインスレッドに投げる
            try:
                future.result()
            except Exception as e:
                executor.shutdown(wait=False, cancel_futures=True)
                raise e
            
            with progress_lock:
                completed_pages += 1
                ui.update_overall(completed_pages, total_tasks, f"全体の進捗 ( {completed_pages} / {total_tasks} ページ完了 )")

    ui.set_determinate(total_tasks, total_tasks, "すべての処理が完了しました")